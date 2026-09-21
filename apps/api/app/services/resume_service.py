"""Resume ingestion: extract text, then extract structured facts from it."""

from __future__ import annotations

import io

import structlog

from app.graph.contracts import ResumeFacts
from app.llm import ChatMessage, LLMError, get_provider
from app.llm.prompts import RESUME_EXTRACTION_PROMPT

log = structlog.get_logger(__name__)

# Enough of a CV for the model to work with; longer documents are almost always
# padded with references and boilerplate.
MAX_EXTRACT_CHARS = 12_000

SUPPORTED_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
    "text/markdown": "txt",
}


class ResumeError(RuntimeError):
    """Raised when a document cannot be read."""


def detect_kind(content_type: str, filename: str) -> str:
    if kind := SUPPORTED_TYPES.get(content_type.split(";")[0].strip()):
        return kind
    # Browsers are inconsistent about content types for .md and .txt.
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix in {"pdf", "docx", "txt", "md"}:
        return "txt" if suffix == "md" else suffix
    raise ResumeError("Unsupported file type. Upload a PDF, DOCX, or plain text file.")


def extract_text(data: bytes, kind: str) -> str:
    """Pull plain text out of an uploaded document."""
    if kind == "txt":
        return data.decode("utf-8", errors="replace")

    if kind == "pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            pages = [page.extract_text() or "" for page in reader.pages]
        except Exception as exc:
            raise ResumeError(f"Could not read that PDF: {exc}") from exc
        return "\n".join(pages)

    if kind == "docx":
        try:
            import docx

            document = docx.Document(io.BytesIO(data))
            blocks = [p.text for p in document.paragraphs]
            # Skills are very often in a table rather than a paragraph.
            for table in document.tables:
                for row in table.rows:
                    blocks.append(" | ".join(cell.text for cell in row.cells))
        except Exception as exc:
            raise ResumeError(f"Could not read that DOCX: {exc}") from exc
        return "\n".join(blocks)

    raise ResumeError("Unsupported file type.")  # pragma: no cover


def normalise(text: str) -> str:
    """Collapse the whitespace noise that PDF extraction always produces."""
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()


async def extract_facts(raw_text: str) -> ResumeFacts:
    """Ask the LLM for structured facts. Degrades to empty rather than failing."""
    if not raw_text.strip():
        return ResumeFacts(
            headline="Empty document", years_experience=0.0, skills=[], projects=[], domains=[]
        )

    provider = get_provider()
    try:
        return await provider.structured(
            system=RESUME_EXTRACTION_PROMPT,
            messages=[ChatMessage(role="user", content=raw_text[:MAX_EXTRACT_CHARS])],
            schema=ResumeFacts,
            max_tokens=2048,
        )
    except LLMError as exc:
        log.warning("resume.extraction_failed", error=str(exc))
        # An upload that cannot be parsed is still usable — the raw text is
        # passed to the interviewer either way.
        return ResumeFacts(
            headline="Could not analyse this CV automatically",
            years_experience=0.0,
            skills=[],
            projects=[],
            domains=[],
        )
